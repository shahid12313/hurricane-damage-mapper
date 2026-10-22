"""Build the project report (docs/Hurricane_Damage_Mapper_Report.docx).

Numbers are read from the training results in runs/ and dashboard/data/, and code excerpts
are taken from the source files, so rerunning this after new training keeps the report current:

    /opt/anaconda3/bin/python3 docs/build_report.py
"""

import ast
import csv
import json
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "docs" / "figures"
OUT = ROOT / "docs" / "Hurricane_Damage_Mapper_Report.docx"

ACCENT = RGBColor(0x1F, 0x4E, 0x8C)
MUTED = RGBColor(0x55, 0x5F, 0x6D)
DAMAGE_LEVELS = ["no-damage", "minor-damage", "major-damage", "destroyed"]
LEVEL_NAMES = ["No damage", "Minor damage", "Major damage", "Destroyed"]


# ---------------------------------------------------------------- results

def read_history(run):
    path = ROOT / "runs" / run / "history.csv"
    if not path.exists():
        return []
    with open(path) as f:
        return [{k: float(v) for k, v in row.items()} for row in csv.DictReader(f)]


def best_row(rows, key):
    return max(rows, key=lambda r: r[key]) if rows else None


def read_confusion(run):
    path = ROOT / "runs" / run / "building_confusion.csv"
    if not path.exists():
        return None
    with open(path) as f:
        return [[int(v) for v in row[1:]] for row in list(csv.reader(f))[1:]]


def hurricane_agreement():
    result = {}
    for path in sorted((ROOT / "dashboard" / "data").glob("hurricane-*.geojson")):
        feats = json.load(open(path))["features"]
        truth = [f["properties"] for f in feats if f["properties"]["source"] == "true"]
        predicted = sum(f["properties"]["source"] == "predicted" for f in feats)
        same = sum(p["model_damage"] == p["damage"] for p in truth)
        result[path.stem] = (predicted, len(truth), same / len(truth))
    return result


loc_hist = read_history("loc_resnet34")
v1_hist = read_history("damage_resnet34")
v2_hist = read_history("damage_v2")
LOC = best_row(loc_hist, "f1")
V1 = best_row(v1_hist, "score")
V2 = best_row(v2_hist, "score")
V3 = best_row(read_history("damage_v3"), "score")
FINAL = V3 or V1
CM1 = read_confusion("damage_resnet34")
CM2 = read_confusion("damage_v2")
CM3 = read_confusion("damage_v3")
CMF = CM3 or CM1
AGREE = hurricane_agreement()
TEST_PATH = ROOT / "runs" / "test_results.json"
TEST = json.load(open(TEST_PATH)) if TEST_PATH.exists() else None
GEO_PATH = ROOT / "outputs" / "georef_check.csv"
GEO = __import__("pandas").read_csv(GEO_PATH) if GEO_PATH.exists() else None


def pct(x, digits=1):
    return f"{100 * x:.{digits}f}%"


def building_accuracy(cm):
    return sum(cm[i][i] for i in range(4)) / sum(map(sum, cm))


def recall(cm):
    """Share of each true damage level classified correctly."""
    return [cm[i][i] / sum(cm[i]) for i in range(4)]


# ---------------------------------------------------------------- figures drawn from the results

def draw_confusion(cm, name, title):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np
    cm = np.array(cm)
    share = cm / cm.sum(1, keepdims=True)
    fig, ax = plt.subplots(figsize=(5.6, 4.6))
    ax.imshow(share, cmap="Blues", vmin=0, vmax=1)
    for i in range(4):
        for j in range(4):
            ax.text(j, i, f"{share[i, j]:.0%}\n({cm[i, j]:,})", ha="center", va="center", fontsize=9,
                    color="white" if share[i, j] > 0.55 else "black")
    short = ["No damage", "Minor", "Major", "Destroyed"]
    ax.set_xticks(range(4), short)
    ax.set_yticks(range(4), short)
    ax.set_xlabel("Predicted level")
    ax.set_ylabel("True level")
    ax.set_title(title)
    fig.savefig(FIG / name, dpi=200, bbox_inches="tight")
    plt.close(fig)


def draw_improvement(runs, v1):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    colors = ["#2ca02c", "#e0a800", "#f28e2b", "#d62728"]
    fig, axes = plt.subplots(1, len(runs), figsize=(6 * len(runs), 4), sharey=True)
    for ax, (title, rows) in zip(axes, runs):
        for key, name, color in zip(DAMAGE_LEVELS, LEVEL_NAMES, colors):
            ax.plot([r["epoch"] for r in rows], [r[f"f1_{key}"] for r in rows], color=color, label=name)
            ax.axhline(v1[f"f1_{key}"], color=color, linestyle=":", linewidth=1)
        ax.set(title=title, xlabel="epoch", ylim=(0, 1))
    axes[0].set_ylabel("validation F1")
    axes[0].legend(fontsize=8)
    fig.savefig(FIG / "improvement.png", dpi=170, bbox_inches="tight")
    plt.close(fig)


draw_confusion(CM1, "confusion_v1.png", "v1: building-level confusion (hold split)")
if CM3:
    draw_confusion(CM3, "confusion_v3.png", "v3: building-level confusion (hold split)")
if v2_hist and read_history("damage_v3"):
    draw_improvement([("v2: minor weight 6, from scratch", v2_hist),
                      ("v3: balanced weights, from v1", read_history("damage_v3"))], V1)


# ---------------------------------------------------------------- code excerpts

def source_of(file, name):
    """Source text of a function or Class.method in a project file."""
    text = (ROOT / file).read_text()
    tree = ast.parse(text)
    parts = name.split(".")
    nodes = tree.body
    node = None
    for part in parts:
        node = next(n for n in nodes if isinstance(n, (ast.FunctionDef, ast.ClassDef)) and n.name == part)
        nodes = node.body
    segment = ast.get_source_segment(text, node)
    indent = " " * node.col_offset
    return "\n".join(line[len(indent):] if line.startswith(indent) else line for line in (indent + segment).splitlines())


# ---------------------------------------------------------------- document helpers

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(section, side, Inches(1))

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(11)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.15
for level, size in ((1, 16), (2, 13), (3, 11.5)):
    st = styles[f"Heading {level}"]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = ACCENT
    st.font.bold = True
    st.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True

counters = {"figure": 0, "table": 0}


def add_runs(par, text):
    """Add text with **bold**, `code` and [[placeholder]] markup."""
    for token in re.split(r"(\*\*.+?\*\*|`.+?`|\[\[.+?\]\])", text):
        if not token:
            continue
        if token.startswith("**"):
            par.add_run(token[2:-2]).bold = True
        elif token.startswith("`"):
            run = par.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif token.startswith("[["):
            run = par.add_run(token[2:-2])
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            par.add_run(token)
    return par


def h(text, level=1):
    return doc.add_heading(text, level)


def p(text, style=None, align=None):
    par = add_runs(doc.add_paragraph(style=style), text)
    if align:
        par.alignment = align
    return par


def bullets(items, numbered=False):
    for item in items:
        p(item, style="List Number" if numbered else "List Bullet")


def shade(element, fill):
    props = element.get_or_add_tcPr() if element.tag.endswith("tc") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def caption(kind, text):
    counters[kind] += 1
    label = "Figure" if kind == "figure" else "Table"
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run(f"{label} {counters[kind]}. ")
    run.bold = True
    run.font.size = Pt(9.5)
    run = par.add_run(text)
    run.italic = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = MUTED
    if kind == "table":
        par.paragraph_format.keep_with_next = True
    return counters[kind]


def figure(name, width, text):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(str(FIG / name), width=Inches(width))
    return caption("figure", text)


def table(headers, rows, widths, text=None, align_right=()):
    if text:
        caption("table", text)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, head in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(head)
        run.bold = True
        run.font.size = Pt(10)
        shade(cell._tc, "DCE6F2")
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            par = add_runs(cells[i].paragraphs[0], str(value))
            for run in par.runs:
                run.font.size = Pt(10)
            if i in align_right:
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
            for par in row.cells[i].paragraphs:
                par.paragraph_format.space_after = Pt(2)
    doc.add_paragraph()
    return t


def code(text, size=8.5):
    for line in text.rstrip().splitlines() or [""]:
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.0
        shade(par._p, "F3F4F6")
        run = par.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def toc():
    par = doc.add_paragraph()
    run = par.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click here and choose Update Field to show the table of contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, placeholder, end):
        run._r.append(el)
    # ask Word to refresh fields (the table of contents) when the file is opened
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def page_numbers():
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(9)


# ================================================================ title page

for _ in range(6):
    doc.add_paragraph()
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Hurricane Damage Mapper")
run.bold = True
run.font.size = Pt(28)
run.font.color.rgb = ACCENT
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run("Automated Building Damage Assessment from Before-and-After Satellite Imagery")
run.font.size = Pt(14)
run.font.color.rgb = MUTED
for _ in range(4):
    doc.add_paragraph()
for line in ["Final Project Report",
             "Shahid Abbas",
             "Master of Science in Computer Science",
             "Department of Computer Science, Lamar University, Beaumont, Texas",
             "Advisor: [[Advisor name]]",
             date.today().strftime("%B %Y")]:
    p(line, align=WD_ALIGN_PARAGRAPH.CENTER)
page_break()

h("Contents", 1)
toc()
p("This report has two parts. **Part A** is the project report: problem, data, method, results and "
  "limitations. **Part B** is a walkthrough of the code, file by file, with a glossary of the terms used.")
page_break()

# ================================================================ abstract

h("Abstract", 1)
p("After a hurricane, emergency managers need to know quickly which buildings are damaged and how badly. "
  "Today this is done largely by field surveys or by analysts inspecting satellite images by hand, which takes days. "
  "This project builds a complete system that does it automatically. Using the public xBD dataset of before-and-after "
  "satellite images with expert damage labels, two deep-learning models were trained: a U-Net that finds buildings in "
  "the pre-disaster image, and a second U-Net that compares the pre- and post-disaster images and labels each building "
  "pixel with one of four damage levels. Building outlines are then extracted, placed on the map, and shown in an "
  "interactive web dashboard.")
p(f"Both models were trained entirely on a laptop (Apple M1 Pro). On the held-out validation split, building "
  f"localization reached a pixel F1 of **{LOC['f1']:.3f}**. The first damage model reached an xView2-style score of "
  f"{V1['score']:.3f} but recognized minor damage poorly (F1 {V1['f1_minor-damage']:.2f}). Fine-tuning it on three times "
  f"more data with balanced class weights raised the score to **{FINAL['score']:.3f}** and roughly doubled "
  f"minor-damage F1 (to {FINAL['f1_minor-damage']:.2f}), at the cost of more false alarms on undamaged buildings; "
  f"overall, {pct(building_accuracy(CMF))} of buildings receive the correct damage level. A failed intermediate attempt "
  "showed how class weights trade neighbouring damage levels off against each other. The project also found that the "
  "georeferencing tags in the xBD GeoTIFF files misplace buildings by a median of 30 metres (up to 80 m, growing with "
  "latitude) across the whole dataset, explained the cause, and corrected it using the label coordinates.")
if TEST:
    p(f"On the untouched test split, the final model reaches a score of **{TEST['damage_v3']['score']:.3f}** "
      f"(building localization F1 {TEST['localization']['f1']:.3f}) and assigns the correct damage level to "
      f"**{pct(TEST['damage_v3']['building_accuracy'])}** of buildings.")

# ================================================================ PART A

page_break()
h("Part A — Project Report", 1)

h("1. Introduction", 1)
h("1.1 Motivation", 2)
p("In August 2017, Hurricane Harvey brought record rainfall to southeast Texas, flooding large parts of Houston and "
  "Beaumont. In the days after a storm like this, responders must decide where to send search teams, where to set up "
  "shelters and how to prioritize aid, and all of these decisions depend on knowing where the damage is. High-resolution "
  "satellite images are usually available within days, but turning thousands of square kilometres of imagery into a "
  "building-by-building damage map is slow when done by hand.")
p("Deep learning offers a way to automate this step. If a model can look at the same neighbourhood before and after a "
  "storm and say which buildings were damaged, the result can be put on a map for responders within hours instead of days.")

h("1.2 Objectives", 2)
bullets([
    "Build a model that finds building footprints in satellite images.",
    "Build a model that classifies each building's damage on the four-level xBD scale: no damage, minor, major, destroyed.",
    "Evaluate both models with the same metrics as the xView2 challenge, so results are comparable with published work.",
    "Present the results on an interactive map that a non-specialist can use.",
    "Do all of this on ordinary hardware (a laptop), without cloud computing.",
], numbered=True)

h("1.3 Contributions", 2)
bullets([
    "A complete, reproducible pipeline from raw xBD GeoTIFFs to a web map, in about 1,000 lines of Python plus one HTML page.",
    "A compact data format that shrinks the dataset from 130 GB to about 9 GB with negligible quality loss, making training on a laptop practical.",
    "Two trained U-Net models with results reported per damage level, per building and per hurricane.",
    "Identification and correction of a georeferencing error in the xBD GeoTIFF metadata.",
    "An interactive dashboard that compares model predictions with expert labels and ranks the hardest-hit areas.",
])

h("2. Background", 1)
h("2.1 The xBD dataset and the xView2 challenge", 2)
p("xBD (Gupta et al., 2019) is the largest public dataset for building damage assessment from satellite imagery. It "
  "contains pairs of high-resolution images taken before and after 19 natural disasters, including hurricanes, floods, "
  "wildfires, earthquakes, tsunamis and volcanic eruptions. Every building is outlined as a polygon, and in the "
  "post-disaster image each building is labelled with a damage level from the Joint Damage Scale.")
table(["Level", "Description (summarized from the xBD Joint Damage Scale)"], [
    ["No damage", "Undisturbed; no sign of water, structural damage, shingle damage or burn marks."],
    ["Minor damage", "Building partially burnt, water surrounding it, volcanic flow nearby, roof elements missing, or visible cracks."],
    ["Major damage", "Partial wall or roof collapse, or encroaching volumes of water or mud."],
    ["Destroyed", "Completely collapsed, scorched, or no longer present."],
], [1.3, 5.2], "Damage levels used in xBD.")
p("The dataset was released with the xView2 challenge, which scored entries on two tasks: finding buildings "
  "(localization) and classifying their damage. This project uses the same scoring (Section 5.5).")

h("2.2 U-Net and transfer learning", 2)
p("U-Net (Ronneberger et al., 2015) is a neural network for image segmentation: it outputs a label for every pixel. "
  "It has an encoder that progressively shrinks the image into abstract features, a decoder that expands them back to "
  "full resolution, and skip connections that pass fine detail from the encoder to the decoder so edges stay sharp. "
  "In this project the encoder is a ResNet-34 (He et al., 2016) pretrained on ImageNet. Starting from a network that "
  "already recognizes edges, textures and shapes (transfer learning) lets the model learn from far fewer examples "
  "than training from scratch.")

h("3. Data", 1)
h("3.1 Splits", 2)
p("The full xBD GeoTIFF release (about 55 GB compressed, 130 GB extracted) contains four splits. Each tile is "
  "1024 × 1024 pixels at roughly 0.5 m per pixel.")
table(["Split", "Image pairs", "Use in this project"], [
    ["tier1", "2,799", "Training (all models)"],
    ["tier3", "6,369", "Extra training data (v2 damage model)"],
    ["hold", "933", "Validation: choosing the best epoch and reporting results"],
    ["test", "933", "Reserved for the final evaluation; also shown on the dashboard"],
], [1.0, 1.2, 4.3], "xBD splits.", align_right=(1,))

h("3.2 Exploration", 2)
p("Before training, the tier1 split was explored in `notebooks/01_explore_xbd.ipynb`. It contains 162,787 labelled "
  "buildings from 10 disasters. Three findings shaped the modelling choices:")
table(["Damage level", "Buildings", "Share"], [
    ["No damage", "117,426", "72.1%"],
    ["Minor damage", "14,980", "9.2%"],
    ["Major damage", "14,161", "8.7%"],
    ["Destroyed", "13,227", "8.1%"],
    ["Un-classified", "2,993", "1.8%"],
], [2.2, 1.4, 1.2], "Damage labels in tier1.", align_right=(1, 2))
bullets([
    "**Strong class imbalance.** Almost three quarters of buildings are undamaged. A model that always answered "
    "\"no damage\" would be 72% accurate and useless, so damaged classes are given more weight in training and results "
    "are reported per class.",
    "**Hurricane damage is mostly flooding.** For Hurricane Harvey, 35.8% of buildings are labelled major damage but "
    "only 1.7% destroyed: roofs look intact and the damage shows as water around buildings (Figure 2). Fire damage, by "
    "contrast, removes buildings entirely.",
    "**Buildings are small.** The median building covers 654 pixels (about 25 × 25), and 5% cover fewer than 87 pixels. "
    "Images therefore cannot be shrunk before training without losing small buildings.",
])
figure("damage_distribution.png", 4.6, "Distribution of damage labels in tier1.")
figure("harvey_pair.jpg", 6.3, "A Hurricane Harvey tile before and after the storm, with labelled building outlines "
       "(orange: major damage). The neighbourhood is flooded but roofs are intact.")

h("4. System Overview", 1)
figure("pipeline.png", 6.5, "The processing pipeline and the file that implements each step.")
bullets([
    "**Preprocess.** Convert the raw GeoTIFFs into compact JPEG images and per-pixel label masks.",
    "**Model 1, localization.** From the pre-disaster image, predict for every pixel whether it belongs to a building.",
    "**Model 2, damage classification.** From the pre- and post-disaster images together, predict every pixel's damage level.",
    "**Building extraction and georeferencing.** Split building pixels into individual buildings, give each one a "
    "damage level, and convert its outline from pixels to longitude and latitude.",
    "**Dashboard.** Show the buildings, coloured by damage, on an interactive satellite map.",
])

h("5. Methodology", 1)
h("5.1 Preprocessing", 2)
p("The GeoTIFF images store pixels as 16-bit integers whose values lie in the 8-bit range 0–255, a format the common "
  "Pillow library cannot open. They are read with `tifffile`, converted to 8-bit colour and saved as JPEG at quality 95. "
  "On a sample of tiles this changed pixel values by 1.46 out of 255 on average, which is invisible, while shrinking "
  "the dataset from 130 GB to about 9 GB and making images much faster to load during training.")
p("Each label file is turned into a mask image of the same size, in which every pixel holds one number:")
table(["Mask value", "Meaning"], [
    ["0", "Background (not a building)"],
    ["1", "No damage"], ["2", "Minor damage"], ["3", "Major damage"], ["4", "Destroyed"],
    ["5", "Un-classified building (ignored when training the damage model)"],
], [1.2, 5.0], "Encoding of the label masks.")
figure("mask_overlay.jpg", 6.0, "A post-disaster tile (left) and the same tile with its generated mask overlaid (right).")

h("5.2 Localization model", 2)
p("The localization model is a U-Net with a ResNet-34 encoder. It takes the 3-channel pre-disaster image and outputs "
  "one value per pixel, the probability that the pixel belongs to a building. The pre-disaster image is used because "
  "buildings are intact and easiest to see before the storm. Training uses the sum of two losses:")
bullets([
    "**Binary cross-entropy**, which scores each pixel independently.",
    "**Dice loss**, 1 − 2|P ∩ T| / (|P| + |T|), which measures the overlap between predicted and true building areas. "
    "Because it is a ratio, small buildings count as much as large stretches of background, which counters the imbalance "
    "between building and background pixels.",
])

h("5.3 Damage classification model", 2)
p("The damage model is also a U-Net with a ResNet-34 encoder, but its input is the pre- and post-disaster images "
  "stacked into a single 6-channel image, so the network can compare the two directly. The pretrained first layer is "
  "adapted to six channels by the `segmentation_models_pytorch` library. The output has five channels: background and "
  "the four damage levels.")
p("Training uses weighted cross-entropy. The weights are 0.5 for background, 1 for no damage and 3 for each damaged "
  "class, so that a mistake on a rare damaged pixel costs more than a mistake on a common undamaged one. Un-classified "
  "buildings are excluded from the loss.")
p("Many top xView2 solutions use a Siamese design, in which the same encoder processes each image separately before "
  "the features are compared. Stacking the images is simpler, reuses the same architecture as the localization model, "
  "and was chosen as a first version.")

h("5.4 Training setup", 2)
table(["Setting", "Value"], [
    ["Hardware", "Apple MacBook Pro, M1 Pro, 16 GB memory (PyTorch MPS backend)"],
    ["Framework", "PyTorch 2.x, segmentation_models_pytorch"],
    ["Training data", "tier1 (2,799 tiles)"],
    ["Input", "Random 512 × 512 crops of the 1024 × 1024 tiles"],
    ["Augmentation", "Random horizontal and vertical flips and 90° rotations"],
    ["Optimizer", "AdamW, learning rate 3 × 10⁻⁴, weight decay 10⁻⁴"],
    ["Learning-rate schedule", "Cosine decay to zero over 30 epochs"],
    ["Batch size", "8"],
    ["Model selection", "Checkpoint with the best validation score (hold split)"],
    ["Training time", "About 3.5 hours per model (about 7 minutes per epoch)"],
], [1.8, 4.7], "Training configuration (both models).")
p("Training is resumable: after every epoch the model, optimizer and schedule are saved, and rerunning the same command "
  "continues from the last finished epoch.")

h("5.5 Evaluation metrics", 2)
p("For a class of interest, **precision** is the fraction of pixels predicted as that class that truly are, "
  "**recall** is the fraction of true pixels that were found, and **F1** is their harmonic mean, "
  "2 × precision × recall / (precision + recall). F1 is high only when both are high. Following the xView2 challenge:")
bullets([
    "**Localization F1:** pixel F1 for building versus background, over the whole validation set.",
    "**Damage F1:** for each of the four damage levels, the pixel F1 on building pixels; then the harmonic mean of the four. "
    "The harmonic mean is pulled down strongly by any weak class, so a model cannot score well by ignoring rare levels.",
    "**Score:** 0.3 × localization F1 + 0.7 × damage F1.",
])
p("Because pixel scores are hard to interpret, results are also reported **per building**: each true building is "
  "assigned the damage level predicted for the majority of its pixels, giving a confusion matrix of true versus "
  "predicted levels.")

h("5.6 From pixels to buildings", 2)
p("The models output per-pixel probabilities, but responders need buildings. For each tile, pixels with building "
  "probability above 0.5 are grouped into connected regions, each treated as one building; regions smaller than 30 "
  "pixels are discarded as noise. Each building's damage level is the one with the highest average probability over "
  "its pixels, and that average is reported as the model's confidence. The outline of each region is traced and "
  "simplified to a polygon.")

h("5.7 Georeferencing", 2)
p("To place buildings on a map, pixel positions must be converted to longitude and latitude. The GeoTIFF files contain "
  "georeferencing tags for this (a tiepoint and a pixel size), and the label files list every building both in pixel "
  "coordinates and in longitude/latitude. When predictions were first placed using the GeoTIFF tags, buildings landed "
  "tens of metres away from the corresponding labelled buildings. The two sources were therefore compared on every "
  "image in the dataset with `src/check_georef.py`.")
if GEO is not None:
    g = GEO.dropna(subset=["offset_median_m"])
    ratio = g.label_ratio / g.expected_ratio
    table(["Measurement (all images with buildings)", "Result"], [
        ["Images checked", f"{len(g):,} of {len(GEO):,}"],
        ["Offset between GeoTIFF tags and labels, median per image", f"{g.offset_median_m.median():.1f} m"],
        ["Images offset by more than 10 m", pct((g.offset_median_m > 10).mean())],
        ["Largest offset", f"{g.offset_max_m.max():.0f} m"],
        ["Labels explained by one affine transform (to 0.1 m)", pct((g.label_fit_max_m < 0.1).mean())],
        ["GeoTIFF ratio of longitude to latitude degrees per pixel", f"{g.tif_ratio.min():.3f} in every image"],
        ["Label ratio relative to square ground pixels, median", f"{ratio.median():.3f}"],
    ], [4.4, 2.0], "Comparison of GeoTIFF georeferencing with label coordinates across the whole dataset.", align_right=(1,))
    by = g.groupby("disaster").agg(lat=("lat", "mean"), off=("offset_median_m", "median")).sort_values("lat", key=abs)
    pick = [d for d in ["palu-tsunami", "hurricane-matthew", "hurricane-harvey", "hurricane-florence", "portugal-wildfire"] if d in by.index]
    table(["Disaster", "Latitude", "Median offset"], [[d, f"{by.loc[d, 'lat']:.1f}°", f"{by.loc[d, 'off']:.1f} m"] for d in pick],
          [2.6, 1.3, 1.5], "The offset grows with distance from the equator.", align_right=(1, 2))
p("The pattern identifies the cause. Every GeoTIFF uses the same pixel size in degrees for longitude and latitude, "
  "which is only correct at the equator: away from it, meridians converge, so a pixel that is square on the ground "
  "spans 1 / cos(latitude) times as many degrees of longitude as of latitude (about 1.15 at Houston). The label "
  "coordinates follow this geometry, the GeoTIFF tags do not, and the error grows with latitude exactly as this "
  "explanation predicts: about 1 m near the equator and almost 50 m at 40°.")
p("As an independent check, a Hurricane Florence tile was placed on a satellite basemap from a different provider "
  "(Esri) both ways (Figure below). With the label coordinates the tile lines up with the basemap; with the GeoTIFF "
  "tags it is visibly stretched and shifted, so every building appears doubled.")
figure("georef_comparison.jpg", 6.5, "The same Hurricane Florence tile (33.6°N) placed on an independent satellite basemap "
       "using the label coordinates (left, aligned) and the GeoTIFF tags (right, shifted and stretched).")
p("The pipeline therefore georeferences each tile from its label coordinates and falls back to the GeoTIFF tags only "
  "for tiles without buildings. A quick search found no published report of this issue for xBD "
  "[[confirm with a search of the xView2 GitHub issues and the xBD paper before calling it a new finding]]; it matters "
  "for anyone who overlays xBD predictions on other map data, such as roads, flood extents or census areas.")

h("5.8 Dashboard", 2)
p("The dashboard is a single web page built with Leaflet, an open-source mapping library. It loads the exported "
  "buildings (GeoJSON) and tile images and provides:")
bullets([
    "a satellite basemap, with the before or after tile image shown on top when zoomed in;",
    "buildings coloured by predicted damage, with a switch to show the expert labels instead;",
    "filters and counts for each damage level;",
    "a popup for every building with its predicted level and confidence, or, for labelled buildings, whether the model agreed;",
    "the share of labelled buildings where the model agreed, per hurricane;",
    "a ranked list of the hardest-hit tiles, which zooms the map to each one when clicked.",
])

h("6. Results", 1)
p("All results in this section and in Section 7 are on the **hold** split, which was used to choose the best training "
  "epoch. They may therefore be slightly optimistic; the final results on the untouched **test** split are in Section 8.")

h("6.1 Building localization", 2)
table(["Metric", "Value"], [
    ["Best epoch", f"{int(LOC['epoch'])} of 30"],
    ["Pixel F1", f"{LOC['f1']:.3f}"],
    ["Precision", f"{LOC['precision']:.3f}"],
    ["Recall", f"{LOC['recall']:.3f}"],
], [2.0, 1.5], "Localization results on the hold split.", align_right=(1,))
figure("loc_curves.png", 6.3, "Localization training loss and validation scores per epoch.")
p("The validation score fluctuated strongly in the first half of training (for example 0.78 at epoch 3 and 0.49 at "
  "epoch 4) while the learning rate was high, then settled between 0.77 and 0.79 from epoch 17 onwards as the cosine "
  "schedule reduced the learning rate. Precision is higher than recall: the model rarely marks non-buildings as "
  "buildings, but misses some building pixels.")
figure("loc_prediction.jpg", 6.3, "Localization on a validation tile from Hurricane Michael.")

h("6.2 Damage classification (first model, v1)", 2)
table(["Metric", "Value"], [
    ["Best epoch", f"{int(V1['epoch'])} of 30"],
    ["Score (0.3 × loc + 0.7 × damage)", f"{V1['score']:.3f}"],
    ["Localization F1", f"{V1['loc_f1']:.3f}"],
    ["Damage F1 (harmonic mean)", f"{V1['damage_f1']:.3f}"],
    ["F1 — no damage", f"{V1['f1_no-damage']:.3f}"],
    ["F1 — minor damage", f"{V1['f1_minor-damage']:.3f}"],
    ["F1 — major damage", f"{V1['f1_major-damage']:.3f}"],
    ["F1 — destroyed", f"{V1['f1_destroyed']:.3f}"],
    ["Building-level accuracy", pct(building_accuracy(CM1))],
], [3.0, 1.5], "First damage model (v1) on the hold split.", align_right=(1,))
p(f"The damage model's own localization F1 ({V1['loc_f1']:.3f}) is higher than the dedicated localization model's "
  f"({LOC['f1']:.3f}), probably because seeing the post-disaster image as well gives it more evidence of where buildings are.")
figure("confusion_v1.png", 4.4, "Building-level confusion matrix of v1. Each row shows how the buildings of one true level were classified.")
r1 = recall(CM1)
p(f"The model is strongest at the two ends of the scale: {pct(r1[0], 0)} of undamaged and {pct(r1[3], 0)} of destroyed "
  f"buildings are classified correctly. Minor damage is the clear weakness: only {pct(r1[1], 0)} of minor-damage "
  f"buildings are recognized, with {pct(CM1[1][2] / sum(CM1[1]), 0)} labelled major and {pct(CM1[1][0] / sum(CM1[1]), 0)} "
  "labelled no damage. This is consistent with the definition of the class: minor damage is often a few missing roof "
  "elements or some water around a building, which is barely visible at 0.5 m per pixel, and it is also the level on "
  "which human labellers disagree most. Section 7 describes how this was improved.")
figure("damage_curves.png", 6.3, "v1 training loss and validation scores per epoch.")
figure("damage_prediction.jpg", 6.5, "v1 damage prediction on a Hurricane Michael validation tile.")

h("6.3 Context", 2)
p("For orientation: Chen et al. (2022) report the official xView2 baseline at an overall score of 0.265, and a single "
  "ResNet-34 version of the first-place Siamese U-Net at 0.717, on the xBD test set; the full first-place solution, an "
  "ensemble of many models, scored higher. The final model here reaches 0.702 on the test split (Section 8) with a "
  "simpler single U-Net trained on a laptop [[check that Chen et al. used the same 933 test tiles before quoting the "
  "comparison as exact]].")

h("7. Improving the Damage Model", 1)
p("The v1 results point to minor damage as the main weakness. Two further versions were trained to address it.")
table(["", "v1", "v2", "v3"], [
    ["Starting point", "ImageNet encoder", "ImageNet encoder", "v1's trained model"],
    ["Training data", "tier1 (2,799)", "tier1 + tier3 (9,168)", "tier1 + tier3 (9,168)"],
    ["Class weights (bg, no, minor, major, destroyed)", "0.5, 1, 3, 3, 3", "0.5, 1, 6, 3, 3", "0.5, 1, 4, 4, 3"],
    ["Learning rate", "3 × 10⁻⁴", "3 × 10⁻⁴", "1 × 10⁻⁴"],
    ["Epochs", "30", "25", "10"],
    ["Training time", "3.5 h", "8 h", "3.3 h"],
], [2.3, 1.3, 1.45, 1.45], "Configuration of the three damage models.")

h("7.1 v2: more data and a higher weight for minor damage", 2)
p("v2 doubled the weight of minor damage, making mistakes on that class twice as costly, and added the tier3 split, "
  "which triples the number of training tiles and adds more variety of buildings and disasters.")
if V2:
    p(f"Minor-damage F1 nearly doubled ({V1['f1_minor-damage']:.2f} → {V2['f1_minor-damage']:.2f}), but major-damage F1 "
      f"collapsed ({V1['f1_major-damage']:.2f} → {V2['f1_major-damage']:.2f}) and the overall score fell from "
      f"{V1['score']:.3f} to {V2['score']:.3f}. The reason is the weighting: minor and major damage are the two levels the "
      "model confuses most, and with minor damage weighted twice as heavily as major, labelling an uncertain building "
      "\"minor\" became the cheaper mistake. The model learned to avoid predicting major damage almost entirely. The "
      "experiment shows that class weights trade neighbouring classes off against each other, rather than improving one "
      "for free.")

h("7.2 v3: balanced weights, fine-tuned from v1", 2)
p("v3 kept the extra tier3 data but gave minor and major damage equal weight (4 each), so neither is the safer guess. "
  "Instead of starting from scratch, it started from v1's trained model, which already separated major damage well, "
  "and trained with a three times lower learning rate so that this knowledge is refined rather than overwritten. "
  "Starting from an already trained model also made it converge in 10 epochs.")
if V3:
    p(f"v3 raised the overall score from {V1['score']:.3f} to **{V3['score']:.3f}**, with clear gains for minor damage "
      f"({V1['f1_minor-damage']:.2f} → {V3['f1_minor-damage']:.2f}), major damage ({V1['f1_major-damage']:.2f} → "
      f"{V3['f1_major-damage']:.2f}) and destroyed buildings ({V1['f1_destroyed']:.2f} → {V3['f1_destroyed']:.2f}), and a "
      f"small loss for no damage ({V1['f1_no-damage']:.2f} → {V3['f1_no-damage']:.2f}). It is the final model used by the dashboard.")

h("7.3 Comparison", 2)
rows = []
for label, key in [("Score", "score"), ("Localization F1", "loc_f1"), ("Damage F1", "damage_f1"),
                   ("F1 — no damage", "f1_no-damage"), ("F1 — minor damage", "f1_minor-damage"),
                   ("F1 — major damage", "f1_major-damage"), ("F1 — destroyed", "f1_destroyed")]:
    rows.append([label] + [f"{run[key]:.3f}" if run else "—" for run in (V1, V2, V3)])
rows.append(["Building-level accuracy"] + [pct(building_accuracy(cm)) if cm else "—" for cm in (CM1, CM2, CM3)])
table(["Metric (hold split, best epoch)", "v1", "v2", "v3"], rows, [2.6, 1.1, 1.1, 1.1],
      "Results of the three damage models.", align_right=(1, 2, 3))
figure("improvement.png", 6.5, "Per-class validation F1 per epoch for v2 (left) and v3 (right). Dotted lines show v1's best values.")
if CM3:
    figure("confusion_v3.png", 4.4, "Building-level confusion matrix of the final model (v3).")
    r3 = recall(CM3)
    p("Compared with v1, the share of buildings recognized correctly changed as follows: "
      + ", ".join(f"{n.lower()} {pct(a, 0)} → {pct(b, 0)}" for n, a, b in zip(LEVEL_NAMES, r1, r3)) + ".")
    p(f"This is a trade-off rather than a pure gain. v3 recognizes far more minor-damage buildings, but it also labels "
      f"{pct(CM3[0][1] / sum(CM3[0]), 0)} of undamaged buildings as minor damage (v1: {pct(CM1[0][1] / sum(CM1[0]), 0)}). "
      f"Because undamaged buildings are the large majority, plain building-level accuracy is about the same "
      f"({pct(building_accuracy(CM1))} for v1, {pct(building_accuracy(CM3))} for v3), while the score, which weights all "
      "damage levels equally, rises sharply. For emergency response this trade-off is usually acceptable: a false alarm "
      "costs a quick second look, while a missed damaged building may mean a household that receives no help. The decision "
      "threshold could also be tuned per use case.")

h("7.4 Final model by hurricane", 2)
p("The dashboard exports cover all hurricane tiles in the hold and test splits. For every labelled building, the "
  "damage level the final model predicted inside its true outline was compared with its label:")
names = {"hurricane-florence": "Florence (Carolinas, 2018)", "hurricane-harvey": "Harvey (Texas, 2017)",
         "hurricane-michael": "Michael (Florida, 2018)", "hurricane-matthew": "Matthew (Haiti, 2016)"}
# agreement of v1 from the first dashboard export (2026-10-19), before it was regenerated with v3
V1_AGREE = {"hurricane-florence": 0.851, "hurricane-harvey": 0.764, "hurricane-michael": 0.587, "hurricane-matthew": 0.362}
rows = [[names.get(k, k), f"{v[1]:,}", pct(V1_AGREE[k]), pct(v[2])]
        for k, v in sorted(AGREE.items(), key=lambda kv: -kv[1][2])]
table(["Hurricane", "Labelled buildings", "v1 agreement", "Final model (v3)"], rows,
      [2.4, 1.4, 1.3, 1.4], "Share of labelled buildings given the same damage level as the experts (hold and test tiles).",
      align_right=(1, 2, 3))
bullets([
    "**Harvey and Matthew improved.** In Haiti (Matthew), most buildings are labelled minor damage, and v3 now recognizes "
    "most of them; v1 had mostly called them major or destroyed.",
    "**Florence and Michael got worse.** These hurricanes have many undamaged buildings, and v3's tendency to call "
    "undamaged buildings minor damage shows most here: on Michael, more than half of the undamaged buildings were "
    "labelled minor damage.",
    "**Matthew still has the fewest buildings found** (about half as many predicted as labelled), because its small, "
    "densely packed houses are often merged into one predicted building.",
])

h("7.5 Dashboard", 2)
figure("dashboard_overview.png", 6.5, "Dashboard overview: Hurricane Harvey tiles in west Houston.")
figure("dashboard_zoom.png", 6.5, "Dashboard zoomed in on a flooded Harvey tile. The after image is placed on the "
       "satellite basemap, and predicted buildings sit on the houses, coloured by damage level.")

h("8. Final Evaluation on the Test Split", 1)
p("Every run used the hold split to choose its best epoch, and the v2 and v3 designs were chosen after looking at hold "
  "results, so hold scores are slightly optimistic. The test split (933 tiles) was never used for any decision. The "
  "localization model, v1 and the final model v3 were each evaluated on it once, with `src/evaluate.py`.")
if TEST:
    L, T1, T3 = TEST["localization"], TEST["damage_v1"], TEST["damage_v3"]
    table(["Metric", "Hold (v3)", "Test (v1)", "Test (v3, final)"], [
        [label, f"{V3[key]:.3f}" if V3 and key in V3 else "—", f"{T1[key]:.3f}", f"{T3[key]:.3f}"]
        for label, key in [("Score", "score"), ("Localization F1", "loc_f1"), ("Damage F1", "damage_f1"),
                           ("F1 — no damage", "f1_no-damage"), ("F1 — minor damage", "f1_minor-damage"),
                           ("F1 — major damage", "f1_major-damage"), ("F1 — destroyed", "f1_destroyed")]
    ] + [["Building-level accuracy", pct(building_accuracy(CM3)) if CM3 else "—",
          pct(T1["building_accuracy"]), pct(T3["building_accuracy"])]],
        [2.4, 1.2, 1.2, 1.5], "Damage models on the test split, with the hold result of v3 for comparison.",
        align_right=(1, 2, 3))
    table(["Metric", "Hold", "Test"], [
        ["Pixel F1", f"{LOC['f1']:.3f}", f"{L['f1']:.3f}"],
        ["Precision", f"{LOC['precision']:.3f}", f"{L['precision']:.3f}"],
        ["Recall", f"{LOC['recall']:.3f}", f"{L['recall']:.3f}"],
    ], [2.0, 1.2, 1.2], "Localization model on the hold and test splits.", align_right=(1, 2))
    gap = T3["score"] - V3["score"]
    p(f"On the test split the final model reaches a score of **{T3['score']:.3f}** "
      f"({'+' if gap >= 0 else '−'}{abs(gap):.3f} compared with hold) and assigns the correct damage level to "
      f"**{pct(T3['building_accuracy'])}** of buildings. The improvement from v1 to v3 holds on unseen data "
      f"({T1['score']:.3f} → {T3['score']:.3f}). These test-split figures are the project's reported results.")
    p("[[Add one or two sentences of your own interpretation of the test results here.]]")
else:
    p("**Status:** the test evaluation has not been run yet (`python -m src.evaluate`).")

h("9. Discussion and Limitations", 1)
bullets([
    "**Minor damage is hard to see.** At about 0.5 m per pixel, the difference between no, minor and major damage is "
    "often a few pixels. Even human labellers disagree most on this level.",
    "**Touching buildings are merged.** Buildings are separated by finding connected regions, so adjacent houses become "
    "one predicted building. Top xView2 solutions split them with additional steps such as watershed segmentation.",
    "**Domain shift.** The training data is mostly US suburbs; performance drops on very different settlements (Haiti).",
    "**Validation-based model choices.** Epochs and model versions were chosen on the hold split; only Section 8 "
    "uses the untouched test split.",
    "**Single runs and combined changes.** Each configuration was trained once, so differences of a few points may "
    "partly be run-to-run noise. v3 changed several things at once (starting point, data, weights, learning rate), so "
    "its improvement cannot be attributed to any single change.",
    "**Simple architecture.** A 6-channel U-Net is simpler than the Siamese networks and ensembles used by top solutions.",
])

h("10. Conclusion and Future Work", 1)
p("This project delivered a complete, working system for building damage assessment after hurricanes: from raw "
  "satellite data, through two trained models, to an interactive map. Trained on a single laptop, it finds buildings "
  f"with a pixel F1 of {LOC['f1']:.3f}, reaches a damage score of {FINAL['score']:.3f} and assigns the correct damage "
  f"level to {pct(building_accuracy(CMF))} of buildings. Minor damage remains the hardest level, but fine-tuning with "
  "balanced weights and more data roughly doubled its F1. Along the way the project identified and corrected a "
  "georeferencing error in the dataset's GeoTIFF metadata.")
p("Future work, roughly in order of expected benefit:")
bullets([
    "A Siamese architecture that processes before and after images with a shared encoder.",
    "Separating touching buildings (for example with watershed segmentation).",
    "Ensembles and test-time augmentation (averaging predictions over flipped and rotated inputs).",
    "Fine-tuning on data from the region of interest to reduce domain shift.",
])

h("References", 1)
for ref in [
    "Gupta, R., Goodman, B., Patel, N., Hosfelt, R., Sajeev, S., Heim, E., Doshi, J., Lucas, K., Choset, H., & Gaston, M. "
    "(2019). Creating xBD: A dataset for assessing building damage from satellite imagery. CVPR Workshops.",
    "Ronneberger, O., Fischer, P., & Brox, T. (2015). U-Net: Convolutional networks for biomedical image segmentation. MICCAI.",
    "He, K., Zhang, X., Ren, S., & Sun, J. (2016). Deep residual learning for image recognition. CVPR.",
    "Chen, H., Nemni, E., Vallecorsa, S., Li, X., Wu, C., & Bromley, L. (2022). Dual-tasks Siamese transformer "
    "framework for building damage assessment. arXiv:2201.10953.",
    "Iakubovskii, P. (2019). Segmentation Models PyTorch. GitHub repository, qubvel/segmentation_models.pytorch.",
    "xView2 challenge: Computer vision for building damage assessment. https://xview2.org",
]:
    p(ref)

# ================================================================ PART B

page_break()
h("Part B — Code Walkthrough", 1)
p("This part explains how the code is organized and what each file does, with the most important functions shown in "
  "full. Code excerpts are copied automatically from the source files when this report is built.")

h("B.1 Project structure", 2)
code("""fyp/
├── README.md               setup, commands and results summary
├── requirements.txt        Python libraries
├── src/                    all Python code
│   ├── xbd.py              reading raw xBD files
│   ├── preprocess.py       raw GeoTIFFs -> JPEG images + masks
│   ├── dataset.py          feeding images to the localization model
│   ├── train.py            localization model: training and scoring
│   ├── damage.py           damage model: data, training and scoring
│   ├── predict.py          both models -> buildings on the map
│   └── progress.py         live training progress in the terminal
├── notebooks/
│   ├── 01_explore_xbd.ipynb             dataset exploration
│   ├── 02_building_localization.ipynb   localization results
│   └── 03_damage_classification.ipynb   damage results
├── dashboard/
│   ├── index.html          the map web page
│   ├── data/               exported buildings (generated)
│   └── images/             tile images for the map (generated)
├── docs/                   this report and its figures
├── data/                   dataset: raw xbd/ and processed/ (not in git)
├── runs/                   trained models and logs (not in git)
└── outputs/                charts and tables from notebook 01""")

h("B.2 Running the project from scratch", 2)
p("All commands run from the project folder with Anaconda's Python (`/opt/anaconda3/bin/python3`, written `python` below):")
table(["Step", "Command", "Time"], [
    ["1. Install libraries", "`python -m pip install -r requirements.txt`", "minutes"],
    ["2. Convert the data", "`python -m src.preprocess`", "~5 min"],
    ["3. Train localization", "`python -m src.train --epochs 30`", "~3.5 h"],
    ["4. Train damage model (v1)", "`python -m src.damage --epochs 30`", "~3.5 h"],
    ["4b. Fine-tune (v3)", "`python -m src.damage --out-dir runs/damage_v3 --train-splits tier1 tier3 --epochs 10 "
     "--lr 1e-4 --init-from runs/damage_resnet34/best.pt --class-weights 0.5 1 4 4 3`", "~3.3 h"],
    ["5. Watch training", "`python -m src.progress`", "—"],
    ["6. Export map data", "`python -m src.predict`", "~25 min"],
    ["7. Open the dashboard", "`python -m http.server 8000 -d dashboard`, then visit localhost:8000", "—"],
    ["8. Rebuild this report", "`python docs/build_report.py`", "seconds"],
], [1.6, 3.9, 0.9], "Commands in order.")

h("B.3 How data flows between files", 2)
table(["File", "Reads", "Writes"], [
    ["src/xbd.py", "data/xbd/ (raw tiles and labels)", "— (library used by the others)"],
    ["src/preprocess.py", "data/xbd/", "data/processed/<split>/images, masks"],
    ["src/train.py", "data/processed/", "runs/loc_resnet34/"],
    ["src/damage.py", "data/processed/", "runs/damage_resnet34/ (v1), damage_v2/, damage_v3/"],
    ["src/predict.py", "data/processed/, data/xbd/ labels, runs/", "dashboard/data/, dashboard/images/"],
    ["dashboard/index.html", "dashboard/data/, dashboard/images/", "— (displays in the browser)"],
], [1.6, 2.5, 2.4], "Inputs and outputs of each file.")

h("B.4 src/xbd.py — reading the raw dataset", 2)
p("**Purpose:** everything needed to read the raw xBD files. Each tile has a pre image, a post image and a JSON label "
  "file that lists every building as a polygon in WKT text format, for example `POLYGON ((807.4 898.1, 760.9 943.3, …))`.")
table(["Function", "What it does"], [
    ["`parse_wkt_polygon`", "Turns a WKT polygon string into an array of (x, y) points."],
    ["`polygon_area`", "Area of a polygon in pixels (shoelace formula)."],
    ["`find_pairs`", "Lists every tile in a split with the paths of its pre image, post image and label."],
    ["`load_image`", "Reads a tile as 8-bit RGB; handles the 16-bit GeoTIFFs that Pillow cannot open."],
    ["`load_buildings`", "Reads all buildings in a label file: id, damage level and outline."],
    ["`building_table`", "One row per building across many tiles (used for the statistics in notebook 01)."],
    ["`show_pair`", "Draws a before/after pair with buildings outlined in their damage colour."],
], [1.7, 4.8])
code(source_of("src/xbd.py", "load_image"))
p("**Why:** the GeoTIFFs store pixels as 16-bit signed integers even though the values only go up to 255. Clipping "
  "and converting to 8-bit gives a normal colour image without any loss.")

h("B.5 src/preprocess.py — compact training data", 2)
p("**Purpose:** convert every tile once into JPEG images and a mask, so training never touches the slow 130 GB of "
  "GeoTIFFs. Tiles are processed in parallel on all CPU cores with `multiprocessing.Pool`, and existing files are "
  "skipped, so the script can be stopped and rerun safely.")
code(source_of("src/preprocess.py", "make_mask"))
p("**How it works:** start with an all-zero (background) image and fill each building polygon with its damage "
  "value using OpenCV's `fillPoly`. The result is the mask format described in Section 5.1.")

h("B.6 src/dataset.py — feeding the localization model", 2)
p("**Purpose:** a PyTorch `Dataset`, the standard interface that PyTorch uses to load training examples. Given an index "
  "it returns one example: the pre-disaster image as a tensor and a mask with 1 for building pixels and 0 elsewhere.")
code(source_of("src/dataset.py", "BuildingDataset.__getitem__"))
bullets([
    "**Random crop:** a random 512 × 512 window, so each epoch sees different parts of each tile and batches fit in memory.",
    "**Augmentation:** random flips and rotations. A building seen from above looks equally valid when rotated, so this "
    "multiplies the variety of training examples for free.",
    "**to_tensor:** scales pixels to 0–1 and normalizes them with the ImageNet mean and standard deviation, because the "
    "pretrained encoder expects inputs in that range.",
    "`cv2.setNumThreads(0)` at the top of the file is required on macOS: without it, OpenCV's own threads deadlock "
    "inside PyTorch's background loading processes and training freezes.",
])

h("B.7 src/train.py — localization training", 2)
p("**Purpose:** build the localization model, train it, score it after every epoch and save checkpoints.")
code(source_of("src/train.py", "build_model"))
p("One line builds the whole U-Net: a ResNet-34 encoder with ImageNet weights, 3 input channels (RGB) and 1 output "
  "channel (building probability, after a sigmoid).")
code(source_of("src/train.py", "dice_loss") + "\n\n\n" + source_of("src/train.py", "loss_fn"))
code(source_of("src/train.py", "evaluate"))
p("**How evaluation works:** over the whole validation set, count true positives (predicted building, really a building), "
  "false positives and false negatives, then compute precision, recall and F1 once at the end. Counting over all pixels "
  "of all tiles, rather than averaging per tile, matches the xView2 definition.")
p("The main function, `train_localization`, follows the standard PyTorch training loop. For each batch it runs "
  "forward (compute predictions), computes the loss, runs backward (compute gradients) and takes an optimizer step "
  "(update the weights). After each epoch it evaluates on the hold split, appends a row to `history.csv`, saves "
  "`best.pt` if the score improved and always saves `last.pt` for resuming.")

h("B.8 src/damage.py — damage training", 2)
p("**Purpose:** the same structure as `train.py`, adapted to damage classification. Its `DamageDataset` returns the pre "
  "and post images stacked into 6 channels and the full mask (values 0–4, with un-classified buildings set to 255 so the "
  "loss ignores them).")
code(source_of("src/damage.py", "build_model"))
code(source_of("src/damage.py", "predict_classes"))
p("**Why split the decision in two:** `argmax` over all five outputs decides building versus background. For the damage "
  "level, only the four damage outputs are compared, so every building pixel always gets a damage level even when the "
  "background output is close.")
code(source_of("src/damage.py", "evaluate"))
p("The function returns the xView2 score, the localization F1, the harmonic-mean damage F1 and the F1 of each damage "
  "level. `building_confusion` computes the building-level confusion matrix used in Section 6.2. The `--class-weights` "
  "command-line option sets the loss weights, `--init-from` starts training from an existing checkpoint, and `--lr` "
  "sets the learning rate; v2 and v3 were trained with these options (Section 7).")

h("B.9 src/predict.py — from models to the map", 2)
p("**Purpose:** run both trained models on hurricane tiles and export everything the dashboard needs.")
table(["Function", "What it does"], [
    ["`load_models`", "Loads the best checkpoint of each model."],
    ["`predict_maps`", "Runs both models on one tile: building probability and damage-level probabilities per pixel."],
    ["`extract_buildings`", "Splits building pixels into buildings and gives each a level, confidence and outline."],
    ["`georeference`", "Builds the pixel → longitude/latitude conversion for a tile (Section 5.7)."],
    ["`true_buildings`", "Reads labelled buildings and records which level the model predicted inside each one."],
    ["`export`", "Loops over all tiles and writes GeoJSON files, tile metadata and small tile images. "
     "`--damage-run` chooses the damage model (default: the final model, runs/damage_v3)."],
], [1.7, 4.8])
code(source_of("src/predict.py", "extract_buildings"))
code(source_of("src/predict.py", "georeference"))
p("**How georeference works:** for every building, the label file gives the same outline in pixels and in "
  "longitude/latitude. Solving a least-squares problem finds the six numbers of the affine transform "
  "(lon, lat) = (x, y, 1) · coef that best maps one to the other. For tiles with no buildings it falls back to the "
  "GeoTIFF tags.")

h("B.10 src/progress.py — watching training", 2)
p("A small terminal display that reads each run's `history.csv` every 30 seconds and shows a progress bar, the "
  "latest and best scores, and an estimated finish time. It only reads files, so it can be started and stopped at any "
  "time without affecting training.")

h("B.10b src/evaluate.py and src/check_georef.py — final checks", 2)
p("`src/evaluate.py` scores the localization model, v1 and v3 once on the test split and writes "
  "`runs/test_results.json` (Section 8). `src/check_georef.py` compares the GeoTIFF georeferencing of every image with "
  "the coordinates in its label file and writes `outputs/georef_check.csv` (Section 5.7); it reads only label files and "
  "image headers, so it runs over all 22,068 images in under a minute.")

h("B.11 dashboard/index.html — the map", 2)
p("A single HTML page with embedded JavaScript, using Leaflet for the map. It needs a local web server because browsers "
  "block pages opened from disk from loading data files.")
table(["Part", "What it does"], [
    ["`init`", "Loads tiles.json, fills the disaster menu and opens a tile directly if the link contains one (#tile_id)."],
    ["`selectDisaster`", "Loads that hurricane's GeoJSON (cached after the first time), draws tile outlines and buildings, updates the statistics."],
    ["`drawBuildings`", "Draws the predicted or true buildings for the checked damage levels, coloured by level, each with a popup."],
    ["`updateImagery`", "When zoomed in, adds the before or after image of each visible tile and removes those out of view."],
    ["`drawAgreement`", "Share of labelled buildings where the model gave the same level."],
    ["`drawPriority`", "Ranks tiles by predicted destroyed (×2) and major-damage buildings."],
], [1.6, 4.9])
p("Buildings are drawn on an HTML canvas (`preferCanvas: true`) rather than as individual SVG elements, which keeps the "
  "map responsive with more than 10,000 buildings on screen.")

h("B.12 Notebooks", 2)
table(["Notebook", "Contents"], [
    ["01_explore_xbd", "Tile and building counts, damage distribution per disaster, building sizes, before/after examples."],
    ["02_building_localization", "Training curves and predicted versus true buildings for the localization model."],
    ["03_damage_classification", "Training curves, building-level confusion matrix, and coloured damage predictions."],
], [2.0, 4.5])

h("B.13 Glossary", 2)
table(["Term", "Meaning"], [
    ["Epoch", "One pass of training over all training tiles."],
    ["Batch", "A group of examples (here 8 crops) processed together in one training step."],
    ["Loss", "A number measuring how wrong the model's predictions are; training adjusts the weights to make it smaller."],
    ["Learning rate", "How big each weight update is. Large values learn fast but unstably; small values are stable but slow."],
    ["Cosine schedule", "Lowers the learning rate smoothly from its starting value to zero over training."],
    ["Checkpoint", "A saved copy of the model weights (best.pt, last.pt)."],
    ["Encoder / decoder", "The parts of a U-Net that shrink the image into features and expand them back to a per-pixel output."],
    ["Pretrained (ImageNet)", "Weights learned on 1.2 million everyday photos, used as a starting point (transfer learning)."],
    ["Augmentation", "Random changes (flips, rotations) to training images that create more variety."],
    ["Class weights", "Multipliers in the loss that make mistakes on rare classes cost more."],
    ["Precision / recall / F1", "Fraction of predictions that are right / fraction of true cases found / their harmonic mean."],
    ["Confusion matrix", "Table of true versus predicted classes; the diagonal holds the correct predictions."],
    ["Validation split", "Data not used for training, used to choose the best epoch (hold)."],
    ["Test split", "Data used only once, for the final reported result (test)."],
    ["MPS", "Metal Performance Shaders: PyTorch's backend for the Apple Silicon GPU."],
    ["GeoTIFF", "A TIFF image with embedded geographic coordinates."],
    ["GeoJSON", "A standard text format for map shapes, used by the dashboard."],
    ["Georeferencing", "Converting image pixel positions to real-world coordinates."],
], [1.7, 4.8])

page_numbers()
doc.save(OUT)
print(f"wrote {OUT}  (final damage model score {FINAL['score']:.3f})")
